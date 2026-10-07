#!/usr/bin/env python3
"""Convert the four source workbooks into the web app's normalized JSON dataset."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "questions.json"
IMAGE_DIR = ROOT / "images" / "questions"

SOURCES = [
    {
        "id": "san-pham-da",
        "title": "Sản phẩm đá",
        "shortTitle": "Sản phẩm đá",
        "description": "Nhận diện mẫu đá, mã sản phẩm và cảm hứng thiết kế VICOSTONE.",
        "path": Path("/Users/hoamai/Downloads/15 câu liên quan đến sản phẩm đá_VCS sưu tầm.xlsx"),
        "sheet": "BM câu hỏi",
        "expected": 15,
    },
    {
        "id": "van-hoa-the-thao",
        "title": "Văn hóa & Thể thao",
        "shortTitle": "Văn hóa & Thể thao",
        "description": "Lịch sử, văn hóa Việt Nam và kiến thức thể thao.",
        "path": Path("/Users/hoamai/Downloads/50 câu liên quan đến văn hóa thể thao_VCS sưu tầm.xlsx"),
        "sheet": "Ngan Hang Cau Hoi 500",
        "expected": 52,
    },
    {
        "id": "vicostone",
        "title": "VICOSTONE",
        "shortTitle": "VICOSTONE",
        "description": "Sản phẩm, công nghệ, thương hiệu và hoạt động của VICOSTONE.",
        "path": Path("/Users/hoamai/Downloads/100 câu liên quan đến VICOSTONE_VCS sưu tầm.xlsx"),
        "sheet": "BM câu hỏi",
        "expected": 100,
    },
    {
        "id": "tap-doan",
        "title": "Tập đoàn Phenikaa",
        "shortTitle": "Tập đoàn Phenikaa",
        "description": "Lịch sử, chiến lược, hệ sinh thái và văn hóa Tập đoàn Phenikaa.",
        "path": Path("/Users/hoamai/Downloads/333 câu liên quan đến tập đoàn_VCS sưu tầm.xlsx"),
        "sheet": "BM câu hỏii",
        "expected": 333,
    },
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def displayed_value(cell: Any) -> str:
    """Preserve the human meaning of cells Excel typed as dates or percentages."""
    value = cell.value
    if isinstance(value, (datetime, date)):
        return f"{value.day}/{value.month}"
    if isinstance(value, (int, float)):
        if "%" in cell.number_format:
            amount = value * 100
            return f"{amount:g}%"
        return f"{value:g}"
    return clean_text(value)


def header_map(ws: Any, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[header_row]:
        if cell.value is not None:
            result[clean_text(cell.value)] = cell.column
    return result


def find_column(headers: dict[str, int], *needles: str) -> int | None:
    for header, column in headers.items():
        if any(needle.casefold() in header.casefold() for needle in needles):
            return column
    return None


def find_header_row(ws: Any) -> int:
    for row in range(1, ws.max_row + 1):
        if any("Nội dung câu hỏi" in clean_text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)):
            return row
    raise ValueError(f"Không tìm thấy hàng tiêu đề trong sheet {ws.title!r}")


def extract_images(ws: Any, set_id: str) -> dict[int, str]:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    images: dict[int, str] = {}
    for image in getattr(ws, "_images", []):
        row = image.anchor._from.row + 1
        extension = (image.format or "png").lower().replace("jpeg", "jpg")
        filename = f"{set_id}-{row:03d}.{extension}"
        (IMAGE_DIR / filename).write_bytes(image._data())
        images[row] = f"/images/questions/{filename}"
    return images


def convert_source(source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if not source["path"].exists():
        raise FileNotFoundError(source["path"])

    workbook = load_workbook(source["path"], data_only=True)
    ws = workbook[source["sheet"]]
    header_row = find_header_row(ws)
    headers = header_map(ws, header_row)

    columns = {
        "number": find_column(headers, "Stt"),
        "department": find_column(headers, "Đơn vị/phòng ban"),
        "category": find_column(headers, "Lĩnh vực"),
        "imageAlt": find_column(headers, "Hình ảnh minh họa"),
        "question": find_column(headers, "Nội dung câu hỏi"),
        "answer": find_column(headers, "Đáp án đúng"),
        "explanation": find_column(headers, "Giải thích/làm rõ đáp án"),
        "difficulty": find_column(headers, "Độ khó", "Phân loại"),
    }
    option_columns = {
        letter: find_column(headers, f"Đáp án {letter}") for letter in "ABCD"
    }

    required = ["number", "department", "category", "question", "answer", "explanation", "difficulty"]
    missing = [name for name in required if columns[name] is None]
    if missing or any(value is None for value in option_columns.values()):
        raise ValueError(f"Thiếu cột trong {source['path'].name}: {missing}")

    image_by_row = extract_images(ws, source["id"])
    questions: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        question = clean_text(ws.cell(row, columns["question"]).value)
        if not question:
            continue

        number = clean_text(ws.cell(row, columns["number"]).value) or str(len(questions) + 1)
        correct = clean_text(ws.cell(row, columns["answer"]).value).upper()
        options = [displayed_value(ws.cell(row, option_columns[letter])) for letter in "ABCD"]
        if correct not in "ABCD" or len(correct) != 1 or any(not option for option in options):
            raise ValueError(f"Câu hỏi không hợp lệ: {source['id']} dòng {row}")

        image_alt = ""
        if columns["imageAlt"]:
            image_alt = clean_text(ws.cell(row, columns["imageAlt"]).value)
            image_alt = re.sub(r"^-+Image of:\s*-*", "Mẫu đá ", image_alt, flags=re.IGNORECASE)

        questions.append(
            {
                "id": f"{source['id']}-{int(float(number)):03d}",
                "setId": source["id"],
                "number": int(float(number)),
                "department": clean_text(ws.cell(row, columns["department"]).value),
                "category": clean_text(ws.cell(row, columns["category"]).value),
                "difficulty": clean_text(ws.cell(row, columns["difficulty"]).value),
                "question": question,
                "options": options,
                "answer": "ABCD".index(correct),
                "explanation": clean_text(ws.cell(row, columns["explanation"]).value),
                "image": image_by_row.get(row),
                "imageAlt": image_alt or None,
            }
        )

    if len(questions) != source["expected"]:
        raise ValueError(
            f"{source['id']}: dự kiến {source['expected']} câu, đọc được {len(questions)}"
        )

    metadata = {key: source[key] for key in ("id", "title", "shortTitle", "description")}
    metadata["questionCount"] = len(questions)
    return metadata, questions


def main() -> None:
    DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)

    sets: list[dict[str, Any]] = []
    questions: list[dict[str, Any]] = []
    for source in SOURCES:
        metadata, converted = convert_source(source)
        sets.append(metadata)
        questions.extend(converted)

    identifiers = [question["id"] for question in questions]
    if len(identifiers) != len(set(identifiers)):
        raise ValueError("ID câu hỏi bị trùng")

    payload = {
        "version": 1,
        "generatedFrom": "4 bộ câu hỏi nội bộ dạng Excel",
        "totalQuestions": len(questions),
        "sets": sets,
        "questions": questions,
    }
    DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Đã tạo {DATA_PATH} với {len(questions)} câu hỏi và {len(list(IMAGE_DIR.glob('*')))} ảnh.")


if __name__ == "__main__":
    main()
